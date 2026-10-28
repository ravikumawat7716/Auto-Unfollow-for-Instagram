#!/usr/bin/env python3
"""
IG Unfollow Bot

Opens a Chrome window, scans who you follow and who follows you, then visits
each account that doesn't follow you back and clicks Following > Unfollow on
its profile, the same way you would by hand. You can watch every step.

Usage (from this folder):
    python ig_unfollow.py scan       Scan your lists and build the Excel report
    python ig_unfollow.py run        Unfollow people who don't follow you back
    python ig_unfollow.py report     Rebuild the Excel report from saved data

Options:
    --limit N    Stop after N unfollows in this run (good for a first test)
    --rescan     Scan again before unfollowing
    --yes        Don't ask for confirmation before starting
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    sys.stdout.reconfigure(errors="replace")  # type: ignore[attr-defined]
except Exception:
    pass

try:
    from playwright.sync_api import Error as PlaywrightError
    from playwright.sync_api import sync_playwright
except ImportError:
    print("Playwright is not installed. Double-click setup.bat first")
    print("(or run: python -m pip install -r requirements.txt)")
    sys.exit(1)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("openpyxl is not installed. Double-click setup.bat first")
    print("(or run: python -m pip install -r requirements.txt)")
    sys.exit(1)


# ---------------------------------------------------------------- paths

IG = "https://www.instagram.com"
FAST = os.environ.get("IGU_FAST") == "1"  # tests only: shorter "thinking" pauses
HEADLESS = os.environ.get("IGU_HEADLESS") == "1"  # tests only

ROOT = DATA_DIR = STATE_FILE = CSV_LOG = REPORT_FILE = SETTINGS_FILE = KEEP_FILE = PROFILE_DIR = Path(".")


def configure_paths(root: Path, profile_dir: Path | None = None) -> None:
    """Point every file the bot uses at `root`."""
    global ROOT, DATA_DIR, STATE_FILE, CSV_LOG, REPORT_FILE, SETTINGS_FILE, KEEP_FILE, PROFILE_DIR
    ROOT = Path(root)
    DATA_DIR = ROOT / "data"
    STATE_FILE = DATA_DIR / "state.json"
    CSV_LOG = DATA_DIR / "actions_log.csv"
    REPORT_FILE = ROOT / "Instagram Unfollow Report.xlsx"
    SETTINGS_FILE = ROOT / "settings.json"
    KEEP_FILE = ROOT / "keep.txt"
    if profile_dir is None:
        # Chrome's own data for the bot window (your Instagram login). Kept outside
        # this folder so it doesn't bloat it.
        base = os.environ.get("LOCALAPPDATA")
        profile_dir = Path(base) / "IGUnfollowBot" / "chrome-profile" if base else ROOT / "chrome-profile"
    PROFILE_DIR = Path(profile_dir)


configure_paths(Path(os.environ.get("IGU_HOME") or Path(__file__).resolve().parent),
                Path(os.environ["IGU_PROFILE"]) if os.environ.get("IGU_PROFILE") else None)


# ---------------------------------------------------------------- settings

DEFAULT_SETTINGS = {
    "delay_seconds": 10,
    "jitter_percent": 30,
    "batch_size": 5,
    "break_min_seconds": 60,
    "break_max_seconds": 120,
    "daily_limit": 150,
    "browse_feed_during_breaks": True,
    "verify_after_each_batch": True,
    "keep_verified": False,
    "my_username": "",
    "browser": "chrome",
}

SETTINGS_HELP = {
    "delay_seconds": "Wait between unfollows. Randomised by jitter_percent.",
    "jitter_percent": "How much to randomise waits, in percent (30 means 7 to 13 s for a 10 s wait).",
    "batch_size": "Unfollows before a longer break.",
    "break_min_seconds": "Shortest break.",
    "break_max_seconds": "Longest break.",
    "daily_limit": "Max unfollows per day. Instagram doesn't publish its limits, so keep this modest.",
    "browse_feed_during_breaks": "Scroll your home feed during breaks, like a person would.",
    "verify_after_each_batch": "After each batch, reopen the last profile to check Instagram really kept the unfollow.",
    "keep_verified": "Never unfollow blue tick accounts.",
    "my_username": "Your Instagram username. Leave empty, the bot finds it by itself.",
    "browser": "chrome, msedge or chromium.",
}


def load_settings() -> dict:
    data: dict = {}
    if SETTINGS_FILE.exists():
        try:
            data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8-sig"))
        except Exception as e:  # noqa: BLE001
            print(f"Could not read settings.json ({e}). Using defaults.")
            data = {}
    settings = dict(DEFAULT_SETTINGS)
    for k, v in data.items():
        if k in DEFAULT_SETTINGS:
            settings[k] = v
    # sanity limits
    settings["delay_seconds"] = max(3, float(settings["delay_seconds"]))
    settings["jitter_percent"] = min(90, max(0, float(settings["jitter_percent"])))
    settings["batch_size"] = max(1, int(settings["batch_size"]))
    settings["break_min_seconds"] = max(5, float(settings["break_min_seconds"]))
    settings["break_max_seconds"] = max(settings["break_min_seconds"], float(settings["break_max_seconds"]))
    settings["daily_limit"] = max(1, int(settings["daily_limit"]))
    if not SETTINGS_FILE.exists():
        write_default_settings()
    return settings


def write_default_settings() -> None:
    out = dict(DEFAULT_SETTINGS)
    out["_help"] = SETTINGS_HELP
    try:
        SETTINGS_FILE.write_text(json.dumps(out, indent=2), encoding="utf-8")
    except OSError:
        pass


# ---------------------------------------------------------------- state

STOP_REASONS = {
    "done": "Finished. Nobody left to unfollow.",
    "daily_limit": "Daily limit reached. Run it again tomorrow and it carries on where it stopped.",
    "run_limit": "Reached the --limit you set for this run.",
    "limited": "Instagram is limiting your actions, so the bot stopped to keep your account safe. Wait at least 24 hours before running again.",
    "auth": "Instagram wants you to log in again or verify your account. Sort that out in the Chrome window, then run again.",
    "errors": "Stopped after 3 problems in a row. Check the Chrome window and the action log, then run again.",
    "stopped": "Stopped by you.",
    "crashed": "Stopped because of an unexpected error. See the action log.",
}

RESULT_LABELS = {
    "unfollowed": "Unfollowed",
    "already_gone": "Already not following",
    "not_found": "Account not available",
    "failed": "Failed",
}

MAX_ATTEMPTS = 3
YES_VALUES = {"yes", "y", "x", "keep", "1", "true", "ok"}
SHEET_NB = "Not following back"


def new_state() -> dict:
    return {
        "version": 1,
        "viewer": None,  # {"id", "username"}
        "profile": None,  # {"followers", "followers_exact", "following", "following_exact"}
        "scanned_at": None,
        "following": [],
        "followers": [],
        "results": {},  # user id -> {"status", "at", "note", "attempts", "run"}
        "keep": [],  # usernames marked "yes" in the Excel Keep column
        "daily": {"date": None, "count": 0},
        "limited_at": None,
        "last_stop": None,
        "log": [],
    }


def load_state() -> dict:
    state = new_state()
    if not STATE_FILE.exists():
        return state
    try:
        state.update(json.loads(STATE_FILE.read_text(encoding="utf-8")))
    except Exception as e:  # noqa: BLE001
        broken = STATE_FILE.with_name("state.broken.json")
        try:
            os.replace(STATE_FILE, broken)
        except OSError:
            pass
        print(f"Saved data could not be read ({e}). Starting fresh. The old file is data/state.broken.json")
        return new_state()
    return state


def save_state(state: dict) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = STATE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, STATE_FILE)


def local_date(ts: float | None = None) -> str:
    return datetime.fromtimestamp(ts if ts is not None else time.time()).strftime("%Y-%m-%d")


def fmt_time(ts) -> str:
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S") if ts else ""


def clear_line() -> None:
    sys.stdout.write("\r" + " " * 70 + "\r")
    sys.stdout.flush()


def log(state: dict, username: str, action: str, result: str, details: str = "") -> None:
    now = time.time()
    state["log"].append({"t": now, "username": username or "", "action": action, "result": result, "details": details or ""})
    if len(state["log"]) > 20000:
        del state["log"][: len(state["log"]) - 20000]
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        new_file = not CSV_LOG.exists()
        with CSV_LOG.open("a", newline="", encoding="utf-8-sig" if new_file else "utf-8") as f:
            w = csv.writer(f)
            if new_file:
                w.writerow(["Time", "Username", "Action", "Result", "Details"])
            w.writerow([fmt_time(now), "@" + username if username else "", action, result, details or ""])
    except OSError:
        pass
    clear_line()
    who = f"@{username} " if username else ""
    print(f"[{datetime.fromtimestamp(now).strftime('%H:%M:%S')}] {who}{action}: {result}" + (f" ({details})" if details else ""))


def set_result(state: dict, user: dict, status: str, note: str, run_id: str) -> None:
    prev = state["results"].get(user["id"]) or {}
    attempts = prev.get("attempts", 0) + (1 if status == "failed" else 0)
    state["results"][user["id"]] = {"status": status, "at": time.time(), "note": note or "", "attempts": attempts, "run": run_id}


# ---------------------------------------------------------------- keep list

def normalize_username(token: str) -> str:
    token = (token or "").strip()
    m = re.search(r"instagram\.com/([^/?#\s]+)", token, re.I)
    if m:
        token = m.group(1)
    token = token.lstrip("@").strip().lower()
    return token if re.fullmatch(r"[a-z0-9._]{1,30}", token) else ""


def read_keep_txt() -> set:
    out: set = set()
    if not KEEP_FILE.exists():
        return out
    for line in KEEP_FILE.read_text(encoding="utf-8-sig").splitlines():
        line = line.split("#", 1)[0]
        for tok in re.split(r"[\s,;]+", line):
            name = normalize_username(tok)
            if name:
                out.add(name)
    return out


def read_keep_from_report() -> set | None:
    """Usernames with yes in the Keep column. None if the report can't be read."""
    if not REPORT_FILE.exists():
        return None
    try:
        wb = load_workbook(REPORT_FILE, read_only=True, data_only=True)
    except PermissionError:
        print("Could not read the Excel report (is it open in Excel?). Using the keep list saved last time.")
        return None
    except Exception:  # noqa: BLE001
        return None
    try:
        if SHEET_NB not in wb.sheetnames:
            return None
        rows = wb[SHEET_NB].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return None
        names = [str(h or "").strip().lower() for h in header]
        keep_i = next((i for i, h in enumerate(names) if h.startswith("keep")), None)
        user_i = names.index("username") if "username" in names else None
        if keep_i is None or user_i is None:
            return None
        keep = set()
        for r in rows:
            if not r or len(r) <= max(keep_i, user_i):
                continue
            if str(r[keep_i] or "").strip().lower() in YES_VALUES:
                u = normalize_username(str(r[user_i] or ""))
                if u:
                    keep.add(u)
        return keep
    finally:
        wb.close()


def sync_keep_from_report(state: dict) -> None:
    keep = read_keep_from_report()
    if keep is not None:
        state["keep"] = sorted(keep)


# ---------------------------------------------------------------- analysis

def analyze(state: dict, settings: dict, keep_txt: set, run_id: str | None = None):
    """Compare the lists. Returns (rows for every account you follow, unfollow queue, counts, following ids)."""
    follower_ids = {u["id"] for u in state["followers"]}
    following_ids = {u["id"] for u in state["following"]}
    keep_excel = set(state.get("keep") or [])
    scanned_at = state.get("scanned_at") or 0
    c = {k: 0 for k in ("mutual", "not_back", "kept", "unfollowed", "already_gone", "failed", "to_unfollow", "you_dont_follow_back")}
    c["following"] = len(state["following"])
    c["followers"] = len(state["followers"])
    rows, queue = [], []
    for u in state["following"]:
        uname = u["username"].lower()
        follows_you = u["id"] in follower_ids
        res = state["results"].get(u["id"]) or {}
        st = res.get("status")
        keep_note = ""
        if follows_you:
            status = "Mutual"
            c["mutual"] += 1
        else:
            c["not_back"] += 1
            if st == "unfollowed" and scanned_at > res.get("at", 0):
                status, keep_note = "Kept (followed again)", "followed again"
                c["kept"] += 1
            elif st in ("unfollowed", "already_gone", "not_found"):
                status = RESULT_LABELS[st]
                c["unfollowed" if st == "unfollowed" else "already_gone"] += 1
            elif uname in keep_excel:
                status, keep_note = "Kept (Excel)", "yes"
                c["kept"] += 1
            elif uname in keep_txt:
                status, keep_note = "Kept (keep.txt)", "keep.txt"
                c["kept"] += 1
            elif settings.get("keep_verified") and u.get("is_verified"):
                status, keep_note = "Kept (verified)", "verified"
                c["kept"] += 1
            elif st == "failed" and (res.get("attempts", 1) >= MAX_ATTEMPTS or (run_id and res.get("run") == run_id)):
                status = "Failed"
                c["failed"] += 1
            else:
                status = "To unfollow"
                c["to_unfollow"] += 1
        row = dict(u)
        row.update(follows_you=follows_you, status=status, keep=keep_note, at=res.get("at"), note=res.get("note", ""))
        rows.append(row)
        if status == "To unfollow":
            queue.append(row)
    c["you_dont_follow_back"] = sum(1 for f in state["followers"] if f["id"] not in following_ids)
    return rows, queue, c, following_ids


def preflight(state: dict, viewer_id: str | None):
    blockers, warnings, notes = [], [], []
    if not state.get("scanned_at"):
        blockers.append("No scan yet. Run a scan first.")
        return blockers, warnings, notes
    viewer = state.get("viewer") or {}
    if viewer_id and viewer.get("id") and str(viewer_id) != str(viewer["id"]):
        blockers.append("You are logged in to a different account than the one that was scanned. Run a scan again.")
    p = state.get("profile") or {}
    got = len(state["followers"])
    fc = p.get("followers")
    if fc:
        need = math.floor(fc * (0.97 if p.get("followers_exact") else 0.9))
        if got < need:
            blockers.append(f"The followers scan looks incomplete ({got} of {fc}). Run a scan again, so nobody who follows you gets unfollowed by mistake.")
        elif not p.get("followers_exact"):
            warnings.append(f"Your profile shows a rounded follower count ({fc}), so the scan ({got}) could only be roughly checked.")
    else:
        warnings.append("Could not read the follower count on your profile, so the scan could not be double-checked.")
    fo = p.get("following")
    if fo and len(state["following"]) < fo * 0.97:
        notes.append(
            f"Your profile says {fo} following and the list has {len(state['following'])}. "
            "The gap is usually deactivated accounts, which Instagram counts but doesn't list. They can't be unfollowed from here."
        )
    if time.time() - state["scanned_at"] > 3 * 86400:
        warnings.append("Your scan is more than 3 days old. Running with --rescan is safer.")
    lim = state.get("limited_at")
    if lim and time.time() - lim < 86400:
        hrs = max(1, round((time.time() - lim) / 3600))
        warnings.append(f"Instagram limited your account about {hrs} hour(s) ago. Waiting at least 24 hours is safer.")
    return blockers, warnings, notes


def apply_scan(state: dict, scan: dict) -> None:
    old = state.get("viewer") or {}
    if str(old.get("id")) != str(scan["viewer"]["id"]):
        fresh = new_state()
        fresh["log"] = state.get("log", [])
        state.clear()
        state.update(fresh)
    state["viewer"] = scan["viewer"]
    state["profile"] = scan["profile"]
    state["following"] = scan["following"]
    state["followers"] = scan["followers"]
    state["scanned_at"] = scan["scanned_at"]
    # Keep the unfollow history. Everything else gets checked again.
    state["results"] = {k: v for k, v in state["results"].items() if v.get("status") == "unfollowed"}


# ---------------------------------------------------------------- Excel report

_report_warned = False


def write_report(state: dict, settings: dict, keep_txt: set) -> bool:
    global _report_warned
    rows, _queue, c, following_ids = analyze(state, settings, keep_txt)
    p = state.get("profile") or {}
    viewer = state.get("viewer") or {}
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EAED")
    keep_fill = PatternFill("solid", fgColor="FFF2CC")
    when = lambda ts: datetime.fromtimestamp(ts) if ts else None  # noqa: E731
    yn = lambda b: "Yes" if b else "No"  # noqa: E731
    link = lambda u: f"{IG}/{u}/"  # noqa: E731

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    today = state["daily"]["count"] if state["daily"].get("date") == local_date() else 0
    last = state.get("last_stop") or {}
    summary = [
        ("Instagram Unfollow Report", ""),
        ("Report updated", datetime.now()),
        ("Account", "@" + viewer["username"] if viewer.get("username") else viewer.get("id", "")),
        ("Last scan", when(state.get("scanned_at"))),
        ("", ""),
        ("Followers (shown on your profile)", p.get("followers") if p.get("followers") is not None else "unknown"),
        ("Following (shown on your profile)", p.get("following") if p.get("following") is not None else "unknown"),
        ("Followers (scanned)", c["followers"]),
        ("Following (scanned)", c["following"]),
        ("", ""),
        ("Follow you back (mutual)", c["mutual"]),
        ("Don't follow you back", c["not_back"]),
        ("    Unfollowed", c["unfollowed"]),
        ("    Kept (Excel, keep.txt, verified, followed again)", c["kept"]),
        ("    Already not following or account gone", c["already_gone"]),
        ("    Failed", c["failed"]),
        ("    Still to unfollow", c["to_unfollow"]),
        ("People who follow you that you don't follow back", c["you_dont_follow_back"]),
        ("", ""),
        ("Unfollowed today", f"{today} of {settings['daily_limit']}"),
        ("Last run ended", f"{fmt_time(last.get('at'))} ({STOP_REASONS.get(last.get('reason'), last.get('reason'))})" if last else ""),
        ("Last time Instagram limited you", when(state.get("limited_at")) or "never"),
        ("", ""),
        ("Settings", ""),
        ("Wait between unfollows (seconds)", f"{settings['delay_seconds']:g} (+/- {settings['jitter_percent']:g}%)"),
        ("Unfollows before a break", settings["batch_size"]),
        ("Break length (seconds)", f"{settings['break_min_seconds']:g} to {settings['break_max_seconds']:g}"),
        ("Max unfollows per day", settings["daily_limit"]),
        ("Keep verified accounts", yn(settings.get("keep_verified"))),
        ("", ""),
        ("How to keep someone", 'Type yes in the Keep column on the "Not following back" sheet, save, close Excel, then run again. Or add their username to keep.txt.'),
    ]
    for r in summary:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws["A"]:
        if cell.value in ("Settings", "How to keep someone"):
            cell.font = bold
    for cell in ws["B"]:
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 70

    def table(title, header, data, widths, link_col=None, keep_col=None):
        sh = wb.create_sheet(title)
        sh.append(header)
        for r in data:
            sh.append(r)
        for i, w in enumerate(widths, start=1):
            sh.column_dimensions[sh.cell(row=1, column=i).column_letter].width = w
        for cell in sh[1]:
            cell.font = bold
            cell.fill = head_fill
        if keep_col:
            sh.cell(row=1, column=keep_col).fill = keep_fill
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
            if link_col:
                cell = row[link_col - 1]
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
        return sh

    table(
        "Following",
        ["Username", "Full name", "Follows you", "Status", "Verified", "Private", "Action time", "Note", "Profile link"],
        [[r["username"], r["full_name"], yn(r["follows_you"]), r["status"], yn(r.get("is_verified")), yn(r.get("is_private")), when(r["at"]), r["note"], link(r["username"])] for r in rows],
        [26, 28, 12, 24, 9, 8, 18, 36, 44],
        link_col=9,
    )
    nb = [r for r in rows if not r["follows_you"]]
    table(
        SHEET_NB,
        ["Keep (type yes)", "Username", "Full name", "Status", "Verified", "Private", "Action time", "Note", "Profile link"],
        [[r["keep"], r["username"], r["full_name"], r["status"], yn(r.get("is_verified")), yn(r.get("is_private")), when(r["at"]), r["note"], link(r["username"])] for r in nb],
        [15, 26, 28, 24, 9, 8, 18, 36, 44],
        link_col=9,
        keep_col=1,
    )
    table(
        "Followers",
        ["Username", "Full name", "You follow them", "Verified", "Profile link"],
        [[u["username"], u["full_name"], yn(u["id"] in following_ids), yn(u.get("is_verified")), link(u["username"])] for u in state["followers"]],
        [26, 28, 15, 9, 44],
        link_col=5,
    )
    table(
        "Action log",
        ["Time", "Username", "Action", "Result", "Details"],
        [[when(e["t"]), "@" + e["username"] if e["username"] else "", e["action"], e["result"], e["details"]] for e in state["log"]],
        [18, 26, 12, 22, 80],
    )

    tmp = REPORT_FILE.with_name("~report.tmp.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, REPORT_FILE)
        _report_warned = False
        return True
    except PermissionError:
        try:
            tmp.unlink()
        except OSError:
            pass
        if not _report_warned:
            clear_line()
            print("  (The Excel report is open in Excel, so it can't be updated right now. Close it and it updates after the next action. The CSV log in data/ is always up to date.)")
            _report_warned = True
        return False


# ---------------------------------------------------------------- browser helpers

OVERLAY_JS = r"""
(() => {
  if (window.top !== window || window.__iguOverlay) return;
  window.__iguOverlay = true;
  const add = () => {
    if (!document.body) { setTimeout(add, 50); return; }
    if (document.getElementById('__igu_status')) return;
    const dot = document.createElement('div');
    dot.id = '__igu_cursor';
    dot.style.cssText = 'position:fixed;left:-50px;top:-50px;width:18px;height:18px;margin:-9px 0 0 -9px;border-radius:50%;' +
      'background:rgba(255,48,64,.6);border:2px solid #fff;box-shadow:0 0 0 1px rgba(0,0,0,.4);z-index:2147483647;pointer-events:none;';
    const bar = document.createElement('div');
    bar.id = '__igu_status';
    bar.style.cssText = 'position:fixed;right:16px;bottom:16px;max-width:360px;padding:10px 14px;border-radius:12px;' +
      'background:rgba(18,18,22,.9);color:#fff;font:13px/1.45 "Segoe UI",system-ui,sans-serif;z-index:2147483647;' +
      'pointer-events:none;white-space:pre-line;box-shadow:0 6px 20px rgba(0,0,0,.35)';
    bar.textContent = window.__iguText || 'IG Unfollow Bot';
    document.body.append(dot, bar);
    addEventListener('mousemove', (e) => { dot.style.left = e.clientX + 'px'; dot.style.top = e.clientY + 'px'; }, true);
    addEventListener('mousedown', () => { dot.style.transform = 'scale(.65)'; }, true);
    addEventListener('mouseup', () => { dot.style.transform = ''; }, true);
  };
  window.__iguSetStatus = (t) => {
    window.__iguText = t;
    const b = document.getElementById('__igu_status');
    if (b) b.textContent = t;
  };
  window.__iguHighlight = (el) => {
    const o = el.style.outline, oo = el.style.outlineOffset;
    el.style.outline = '3px solid #ff3040';
    el.style.outlineOffset = '2px';
    setTimeout(() => { el.style.outline = o; el.style.outlineOffset = oo; }, 1500);
  };
  add();
})();
"""

JS_COMMON = r"""
const vis = (el) => { const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
const T = (el) => ((el && (el.innerText || el.textContent)) || '').trim();
const followScope = () => document.querySelector('header') ? 'header button, header [role="button"]' : 'main button, main [role="button"]';
"""

JS_PAGE_INFO = "() => {" + JS_COMMON + r"""
  const NAMES = ['Following', 'Follow', 'Follow Back', 'Requested'];
  let btn = null;
  for (const el of document.querySelectorAll(followScope())) {
    const t = T(el);
    if (NAMES.includes(t) && vis(el)) { btn = t; break; }
  }
  const dialogs = [...document.querySelectorAll('[role="dialog"]')].filter(vis).map((d) => T(d).slice(0, 500));
  const body = T(document.body).slice(0, 5000);
  return {
    url: location.href,
    btn,
    dialogs,
    unavailable: !btn && /Sorry, this page isn't available|Profile isn't available|This account isn't available|User not found/i.test(body),
  };
}"""

# Finds a clickable element, outlines it and returns its position.
# kind "following" = the Following button on a profile; anything else = a button with that exact text inside a dialog.
JS_TARGET = "(kind) => {" + JS_COMMON + r"""
  let el = null;
  if (kind === 'following') {
    el = [...document.querySelectorAll(followScope())].find((e) => T(e) === 'Following' && vis(e)) || null;
  } else {
    const ds = [...document.querySelectorAll('[role="dialog"]')].filter(vis).reverse();
    for (const d of ds) {
      el = [...d.querySelectorAll('button, [role="button"]')].find((e) => T(e) === kind && vis(e) && !e.dataset.iguClicked) || null;
      if (el) break;
    }
  }
  if (!el) return null;
  let r = el.getBoundingClientRect();
  if (r.top < 0 || r.bottom > innerHeight || r.left < 0 || r.right > innerWidth) {
    el.scrollIntoView({ block: 'center', inline: 'center' });
    r = el.getBoundingClientRect();
  }
  el.dataset.iguClicked = '1';
  if (window.__iguHighlight) window.__iguHighlight(el);
  return { x: r.left, y: r.top, w: r.width, h: r.height };
}"""

JS_MY_USERNAME = r"""() => {
  for (const img of document.querySelectorAll('a[href] img[alt$="profile picture"]')) {
    const a = img.closest('a');
    if (!a || a.closest('main, header, article, [role="dialog"]')) continue;
    const m = (a.getAttribute('href') || '').match(/^\/([A-Za-z0-9._]{1,30})\/?$/);
    if (m) return m[1];
  }
  return null;
}"""

JS_COUNTS = r"""() => {
  const h = document.querySelector('header');
  if (!h) return null;
  const txt = h.innerText || '';
  const g = (re) => { const m = txt.match(re); return m ? m[1] : null; };
  return { followers: g(/([\d.,]+\s*[KkMm]?)\s+followers?\b/), following: g(/([\d.,]+\s*[KkMm]?)\s+following\b/) };
}"""

JS_LIST_PAGE = r"""async ({ kind, id, maxId }) => {
  const ck = (n) => { const p = ('; ' + document.cookie).split('; ' + n + '='); return p.length > 1 ? decodeURIComponent(p.pop().split(';')[0]) : ''; };
  let url = '/api/v1/friendships/' + id + '/' + kind + '/?count=50';
  if (kind === 'followers') url += '&search_surface=follow_list_page';
  if (maxId) url += '&max_id=' + encodeURIComponent(maxId);
  let appId = '936619743392459';
  try { appId = localStorage.getItem('instagramWebFBAppId') || appId; } catch (e) {}
  const headers = { 'x-ig-app-id': appId, 'x-requested-with': 'XMLHttpRequest', 'x-csrftoken': ck('csrftoken') };
  try { const claim = sessionStorage.getItem('www-claim-v2'); if (claim) headers['x-ig-www-claim'] = claim; } catch (e) {}
  let r;
  try { r = await fetch(url, { headers, credentials: 'include' }); } catch (e) { return { status: 0, error: String(e) }; }
  const text = await r.text();
  let j = null;
  try { j = JSON.parse(text); } catch (e) {}
  if (!j) return { status: r.status, error: text.slice(0, 120) };
  return {
    status: r.status,
    message: String(j.message || ''),
    next: j.next_max_id == null ? null : String(j.next_max_id),
    users: (j.users || []).map((u) => ({
      id: String(u.pk ?? u.pk_id ?? u.id),
      username: String(u.username || ''),
      full_name: String(u.full_name || ''),
      is_verified: !!u.is_verified,
      is_private: !!u.is_private,
    })),
  };
}"""

BLOCK_RE = re.compile(
    r"try again later|action blocked|we restrict certain activity|we limit how often|to protect our community|"
    r"temporarily blocked|unusual activity|automated behavio",
    re.I,
)
AUTH_RE = re.compile(r"/(challenge|accounts/login|accounts/suspended|accounts/disabled)", re.I)


class ScanError(Exception):
    def __init__(self, message: str, kind: str = "error"):
        super().__init__(message)
        self.kind = kind


def is_closed_error(e: Exception) -> bool:
    s = str(e).lower()
    return any(x in s for x in ("has been closed", "target closed", "browser has been closed", "connection closed", "browser closed"))


def safe_eval(page, js: str, arg=None):
    try:
        return page.evaluate(js) if arg is None else page.evaluate(js, arg)
    except PlaywrightError as e:
        if is_closed_error(e):
            raise
        return None


def set_status(page, text: str) -> None:
    safe_eval(page, "(t) => window.__iguSetStatus && window.__iguSetStatus(t)", "IG Unfollow Bot\n" + text)


def pause(page, lo: float, hi: float) -> None:
    page.wait_for_timeout(int(random.uniform(lo, hi) * (0.15 if FAST else 1.0) * 1000))


def goto(page, url: str) -> bool:
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=45000)
        return True
    except PlaywrightError as e:
        if is_closed_error(e):
            raise
        return False


def human_click(page, box: dict) -> None:
    """Move the mouse over the element like a hand would, pause, then click."""
    x = box["x"] + box["w"] * random.uniform(0.3, 0.7)
    y = box["y"] + box["h"] * random.uniform(0.3, 0.7)
    page.mouse.move(x, y, steps=random.randint(14, 30))
    pause(page, 0.25, 0.7)
    page.mouse.down()
    page.wait_for_timeout(random.randint(60, 140))
    page.mouse.up()


def press_escape(page) -> None:
    try:
        page.keyboard.press("Escape")
    except PlaywrightError as e:
        if is_closed_error(e):
            raise


def check_blocked(info: dict | None):
    """('auth' | 'limited', note) when Instagram is pushing back, else None."""
    if not info:
        return None
    m = AUTH_RE.search(info.get("url") or "")
    if m:
        return "auth", "Instagram opened its " + m.group(1) + " page"
    for d in info.get("dialogs") or []:
        if BLOCK_RE.search(d):
            return "limited", "Instagram said: " + " ".join(d.split())[:160]
    return None


def page_info(page):
    return safe_eval(page, JS_PAGE_INFO)


def wait_profile_ready(page, timeout: float = 20.0):
    end = time.time() + timeout
    info = None
    while time.time() < end:
        info = page_info(page)
        if info and (info["btn"] or info["unavailable"] or check_blocked(info)):
            return info
        page.wait_for_timeout(300)
    return info


def dismiss_popups(page) -> None:
    info = page_info(page)
    if info and "/accounts/onetap" in info["url"]:
        goto(page, IG + "/")
        page.wait_for_timeout(1500)
    for label in ("Not Now", "Not now"):
        box = safe_eval(page, JS_TARGET, label)
        if box:
            human_click(page, box)
            page.wait_for_timeout(600)


def move_to_middle(page) -> None:
    size = safe_eval(page, "() => [innerWidth, innerHeight]") or [1200, 800]
    page.mouse.move(size[0] * random.uniform(0.4, 0.6), size[1] * random.uniform(0.4, 0.6), steps=random.randint(10, 20))


def countdown(page, seconds: float, label: str, status: str, scroll: bool = False) -> None:
    end = time.time() + max(0.0, seconds)
    next_scroll = time.time() + random.uniform(2, 4)
    try:
        while True:
            left = end - time.time()
            if left <= 0:
                break
            msg = f"{label}: {int(math.ceil(left))}s"
            set_status(page, f"{msg}\n{status}")
            sys.stdout.write("\r  " + msg + " " * 12)
            sys.stdout.flush()
            if scroll and time.time() >= next_scroll:
                dy = random.choice([1, 1, 1, -1]) * random.randint(250, 750)
                try:
                    page.mouse.wheel(0, dy)
                except PlaywrightError as e:
                    if is_closed_error(e):
                        raise
                next_scroll = time.time() + random.uniform(2.5, 6)
            page.wait_for_timeout(int(min(1.0, left) * 1000))
    finally:
        clear_line()


def jitter(seconds: float, pct: float) -> float:
    return max(1.0, seconds * (1 + random.uniform(-pct, pct) / 100))


# ---------------------------------------------------------------- login and scan

def cookies(ctx) -> dict:
    return {c["name"]: c["value"] for c in ctx.cookies(IG)}


def viewer_id(ctx) -> str | None:
    return cookies(ctx).get("ds_user_id") or None


def logged_in(ctx) -> bool:
    ck = cookies(ctx)
    return bool(ck.get("sessionid")) and bool(ck.get("ds_user_id"))


def ensure_logged_in(page, ctx) -> None:
    goto(page, IG + "/")
    if logged_in(ctx):
        page.wait_for_timeout(1500)
        dismiss_popups(page)
        return
    print()
    print("Log in to Instagram in the Chrome window that just opened.")
    print("You type your password into Instagram yourself. This script never sees it.")
    print("Waiting for you (up to 15 minutes)...")
    end = time.time() + 900
    while time.time() < end:
        set_status(page, "Log in to Instagram in this window.\nThe bot waits for you.")
        page.wait_for_timeout(2000)
        if logged_in(ctx):
            print("Logged in. This window remembers you next time.")
            page.wait_for_timeout(3000)
            goto(page, IG + "/")
            page.wait_for_timeout(2000)
            dismiss_popups(page)
            return
    raise SystemExit("Timed out waiting for you to log in.")


def parse_count(text: str | None):
    """'2,981' -> (2981, True). '1.2K' -> (1200, False)."""
    if not text:
        return None, False
    t = text.strip().replace(" ", "")
    mult = 1
    if t[-1:] in "Kk":
        mult, t = 1000, t[:-1]
    elif t[-1:] in "Mm":
        mult, t = 1_000_000, t[:-1]
    try:
        if mult == 1:
            return int(t.replace(",", "").replace(".", "")), True
        return int(float(t.replace(",", "")) * mult), False
    except ValueError:
        return None, False


def find_my_username(page, settings: dict) -> str:
    end = time.time() + 15
    while time.time() < end:
        name = safe_eval(page, JS_MY_USERNAME)
        if name:
            return name
        page.wait_for_timeout(500)
    if settings.get("my_username"):
        return normalize_username(settings["my_username"])
    while True:
        name = normalize_username(input("Could not find your username on the page. Type your Instagram username: "))
        if name:
            return name


def read_profile_counts(page, username: str) -> dict:
    goto(page, f"{IG}/{username}/")
    end = time.time() + 20
    raw = None
    while time.time() < end:
        raw = safe_eval(page, JS_COUNTS)
        if raw and (raw.get("followers") or raw.get("following")):
            break
        page.wait_for_timeout(500)
    raw = raw or {}
    fc, fe = parse_count(raw.get("followers"))
    gc, ge = parse_count(raw.get("following"))
    return {"followers": fc, "followers_exact": fe, "following": gc, "following_exact": ge}


def fetch_all(page, kind: str, vid: str, total: int | None) -> list:
    out, seen, max_id = [], set(), None
    for _ in range(5000):
        res = None
        for attempt in range(4):
            res = safe_eval(page, JS_LIST_PAGE, {"kind": kind, "id": vid, "maxId": max_id})
            if res and res.get("status") == 200 and "users" in res:
                break
            st = res.get("status") if res else "no response"
            msg = (res.get("message") or res.get("error") or "") if res else ""
            if re.search(r"checkpoint|challenge|login_required", msg, re.I):
                raise ScanError(f"Instagram wants you to verify your account ({msg}).", "auth")
            if re.search(r"feedback_required|spam", msg, re.I):
                raise ScanError(f"Instagram is limiting your account ({msg}).", "limited")
            if attempt == 3:
                raise ScanError(f"Could not load your {kind} list (HTTP {st}{': ' + msg if msg else ''}).")
            wait = 60 * (attempt + 1) if st == 429 else 5 * (attempt + 1)
            clear_line()
            print(f"  Instagram answered HTTP {st} while reading your {kind}. Waiting {wait}s and trying again...")
            countdown(page, wait, "Waiting before retry", f"Scanning {kind}")
        for u in res["users"]:
            if u["id"] not in seen:
                seen.add(u["id"])
                out.append(u)
        progress = f"Scanning {kind}: {len(out)}" + (f" of {total}" if total else "")
        set_status(page, progress)
        sys.stdout.write("\r  " + progress + " " * 10)
        sys.stdout.flush()
        nxt = res.get("next")
        if not nxt or not res["users"] or nxt == max_id:
            break
        max_id = nxt
        page.wait_for_timeout(int(random.uniform(1.2, 3.0) * (0.1 if FAST else 1.0) * 1000))
    clear_line()
    print(f"  {kind.capitalize()}: {len(out)}" + (f" (your profile shows {total})" if total else ""))
    return out


def do_scan(page, ctx, state: dict, settings: dict, keep_txt: set) -> bool:
    vid = viewer_id(ctx)
    print()
    print("Scanning your lists...")
    username = find_my_username(page, settings)
    profile = read_profile_counts(page, username)
    print(f"  Account: @{username}")
    set_status(page, f"Scanning @{username}")
    try:
        following = fetch_all(page, "following", vid, profile.get("following"))
        pause(page, 2, 4)
        followers = fetch_all(page, "followers", vid, profile.get("followers"))
    except ScanError as e:
        if e.kind in ("limited", "auth"):
            state["limited_at"] = time.time()
        log(state, "", "Scan", "Failed", str(e))
        save_state(state)
        write_report(state, settings, keep_txt)
        print("Scan failed: " + str(e))
        return False
    apply_scan(state, {
        "viewer": {"id": str(vid), "username": username},
        "profile": profile,
        "following": following,
        "followers": followers,
        "scanned_at": time.time(),
    })
    _rows, _q, c, _ids = analyze(state, settings, keep_txt)
    log(state, "", "Scan", "Done", f"{c['following']} following, {c['followers']} followers, {c['not_back']} don't follow back")
    save_state(state)
    write_report(state, settings, keep_txt)
    set_status(page, f"Scan done. {c['not_back']} accounts don't follow you back.")
    print()
    print(f"  Mutual: {c['mutual']}   Don't follow you back: {c['not_back']}   Left to unfollow: {c['to_unfollow']}")
    return True


# ---------------------------------------------------------------- unfollow by clicking

def unfollow_via_ui(page, username: str, status: str):
    """Open the profile and click Following > Unfollow. Returns (kind, note)."""
    set_status(page, f"Opening @{username}\n{status}")
    if not goto(page, f"{IG}/{username}/"):
        return "failed", "Profile page did not load"
    info = wait_profile_ready(page)
    blocked = check_blocked(info)
    if blocked:
        return blocked
    if not info:
        return "failed", "Profile page did not load"
    if info["unavailable"]:
        return "not_found", "Account not available (deleted or deactivated)"
    btn = info["btn"]
    if btn is None:
        return "failed", "Could not find the Follow button on the profile"
    if btn == "Follow Back":
        return "already_gone", "Already not following (they follow you)"
    if btn == "Follow":
        return "already_gone", "Already not following"
    if btn == "Requested":
        return "already_gone", "Only a pending follow request"

    set_status(page, f"Unfollowing @{username}\n{status}")
    pause(page, 1.2, 3.0)  # look at the profile for a moment
    box = safe_eval(page, JS_TARGET, "following")
    if not box:
        return "failed", "The Following button disappeared"
    human_click(page, box)

    menu_item = None
    end = time.time() + 8
    while time.time() < end:
        page.wait_for_timeout(250)
        blocked = check_blocked(page_info(page))
        if blocked:
            return blocked
        menu_item = safe_eval(page, JS_TARGET, "Unfollow")
        if menu_item:
            break
    if not menu_item:
        press_escape(page)
        return "failed", "The Unfollow option did not appear"
    pause(page, 0.4, 1.2)
    human_click(page, menu_item)

    clicks = 1
    end = time.time() + 12
    while time.time() < end:
        page.wait_for_timeout(400)
        info = page_info(page)
        blocked = check_blocked(info)
        if blocked:
            return blocked
        if info and info["btn"] in ("Follow", "Follow Back"):
            pause(page, 0.4, 1.0)
            return "unfollowed", "Note: this account follows you" if info["btn"] == "Follow Back" else ""
        if clicks < 2:
            confirm = safe_eval(page, JS_TARGET, "Unfollow")  # private accounts ask to confirm
            if confirm:
                pause(page, 0.4, 1.0)
                human_click(page, confirm)
                clicks += 1
    press_escape(page)
    return "failed", "Clicked Unfollow but the button still shows Following"


def verify_unfollowed(page, username: str, status: str):
    """Reopen a profile we just unfollowed. 'reverted' means Instagram quietly undid it."""
    set_status(page, f"Checking @{username} stayed unfollowed\n{status}")
    if not goto(page, f"{IG}/{username}/"):
        return "unknown", ""
    info = wait_profile_ready(page)
    blocked = check_blocked(info)
    if blocked:
        return blocked
    if info and info["btn"] == "Following":
        return "reverted", "Instagram undid the unfollow"
    return "ok", ""


def take_break(page, seconds: float, status: str, settings: dict) -> None:
    browse = settings.get("browse_feed_during_breaks", True)
    if browse:
        goto(page, IG + "/")
        page.wait_for_timeout(1500)
        dismiss_popups(page)
        move_to_middle(page)
    countdown(page, seconds, "Break, browsing your feed" if browse else "Break", status, scroll=browse)


# ---------------------------------------------------------------- the run

def ask_yes(question: str, assume_yes: bool) -> bool:
    if assume_yes:
        return True
    try:
        return input(question + " [y/N]: ").strip().lower() in ("y", "yes")
    except EOFError:
        return False


def do_run(page, ctx, state: dict, settings: dict, keep_txt: set, args) -> str:
    run_id = str(int(time.time()))
    blockers, warnings, notes = preflight(state, viewer_id(ctx))
    for n in notes:
        print("Note: " + n)
    if blockers:
        for b in blockers:
            print("Can't start: " + b)
        return "blocked"
    if warnings:
        for w in warnings:
            print("Warning: " + w)
        if not ask_yes("Start anyway?", args.yes):
            return "stopped"

    s = settings
    run_count = in_batch = errors_in_row = 0
    last_unfollowed = None
    reason = "stopped"
    limit_txt = f", this run: {args.limit}" if args.limit else ""
    print()
    print(f"Starting. Daily limit {s['daily_limit']}{limit_txt}. Press Ctrl+C here, or close the Chrome window, to stop.")
    log(state, "", "Run", "Started", f"daily limit {s['daily_limit']}{limit_txt}")
    try:
        while True:
            today = local_date()
            if state["daily"].get("date") != today:
                state["daily"] = {"date": today, "count": 0}
            if state["daily"]["count"] >= s["daily_limit"]:
                reason = "daily_limit"
                break
            if args.limit and run_count >= args.limit:
                reason = "run_limit"
                break
            _rows, queue, c, _ids = analyze(state, s, keep_txt, run_id)
            if not queue:
                reason = "done"
                break
            u = queue[0]
            status = f"Today {state['daily']['count']} of {s['daily_limit']}  |  Left {c['to_unfollow']}"

            kind, note = unfollow_via_ui(page, u["username"], status)
            acted = False
            if kind == "unfollowed":
                errors_in_row = 0
                in_batch += 1
                run_count += 1
                acted = True
                state["daily"]["count"] += 1
                last_unfollowed = u
                set_result(state, u, "unfollowed", note, run_id)
                log(state, u["username"], "Unfollow", "Done", note)
            elif kind in ("already_gone", "not_found"):
                errors_in_row = 0
                set_result(state, u, kind, note, run_id)
                log(state, u["username"], "Check", "Skipped", note)
            elif kind in ("limited", "auth"):
                state["limited_at"] = time.time()
                log(state, u["username"], "Unfollow", "Stopped by Instagram", note)
                reason = kind
                break
            else:
                errors_in_row += 1
                set_result(state, u, "failed", note, run_id)
                log(state, u["username"], "Unfollow", "Failed", note)
            save_state(state)
            write_report(state, s, keep_txt)

            if errors_in_row >= 3:
                reason = "errors"
                break
            if state["daily"]["count"] >= s["daily_limit"] or (args.limit and run_count >= args.limit) or len(queue) <= 1:
                continue
            status = f"Today {state['daily']['count']} of {s['daily_limit']}  |  Left {max(0, c['to_unfollow'] - 1)}"

            if acted and in_batch >= s["batch_size"]:
                in_batch = 0
                if s.get("verify_after_each_batch", True) and last_unfollowed:
                    pause(page, 1.0, 2.0)
                    v, vnote = verify_unfollowed(page, last_unfollowed["username"], status)
                    if v in ("limited", "auth"):
                        state["limited_at"] = time.time()
                        log(state, last_unfollowed["username"], "Verify", "Stopped by Instagram", vnote)
                        reason = v
                        break
                    if v == "reverted":
                        state["limited_at"] = time.time()
                        set_result(state, last_unfollowed, "failed", vnote, run_id)
                        log(state, last_unfollowed["username"], "Verify", "Instagram undid it", "Unfollows are being reversed, which means Instagram is limiting you")
                        reason = "limited"
                        break
                    log(state, last_unfollowed["username"], "Verify", "OK", "Still unfollowed after reload")
                take_break(page, random.uniform(s["break_min_seconds"], s["break_max_seconds"]), status, s)
            elif acted:
                countdown(page, jitter(s["delay_seconds"], s["jitter_percent"]), "Next unfollow", status)
            else:
                countdown(page, random.uniform(2, 5), "Next account", status)
    except KeyboardInterrupt:
        reason = "stopped"
    except PlaywrightError as e:
        if is_closed_error(e):
            reason = "stopped"
            log(state, "", "Run", "Chrome closed", "The Chrome window was closed")
        else:
            reason = "crashed"
            log(state, "", "Run", "Error", str(e).splitlines()[0][:200])
    except Exception as e:  # noqa: BLE001
        reason = "crashed"
        log(state, "", "Run", "Error", f"{type(e).__name__}: {e}"[:200])
    finally:
        state["last_stop"] = {"reason": reason, "at": time.time()}
        log(state, "", "Run", "Stopped", STOP_REASONS.get(reason, reason))
        save_state(state)
        write_report(state, s, keep_txt)
    try:
        set_status(page, STOP_REASONS.get(reason, reason))
    except Exception:  # noqa: BLE001
        pass
    return reason


# ---------------------------------------------------------------- main

def launch_browser(pw, settings: dict):
    PROFILE_DIR.mkdir(parents=True, exist_ok=True)
    common: dict = {"user_data_dir": str(PROFILE_DIR), "headless": HEADLESS, "locale": "en-US"}
    if HEADLESS:
        common["viewport"] = {"width": 1280, "height": 900}
    else:
        common["no_viewport"] = True
        common["args"] = ["--start-maximized"]
    wanted = str(settings.get("browser") or "chrome").lower()
    order = [wanted] + [b for b in ("chrome", "msedge", "chromium") if b != wanted]
    names = {"chrome": "Google Chrome", "msedge": "Microsoft Edge", "chromium": "Chromium"}
    problems = []
    for b in order:
        kw = dict(common)
        if b != "chromium":
            kw["channel"] = b
        try:
            ctx = pw.chromium.launch_persistent_context(**kw)
            print(f"Opened {names.get(b, b)}.")
            return ctx
        except PlaywrightError as e:
            first = str(e).strip().splitlines()[0] if str(e).strip() else repr(e)
            low = str(e).lower()
            if "already in use" in low or "singletonlock" in low or "profile appears to be in use" in low:
                raise SystemExit("The bot's Chrome window is already open from another run. Close it first.")
            problems.append(f"  {names.get(b, b)}: {first[:160]}")
    raise SystemExit("Could not open a browser.\n" + "\n".join(problems) +
                     "\nInstall Google Chrome, or run: python -m playwright install chromium")


def session(ctx, args, state: dict, settings: dict, keep_txt: set) -> int:
    """Everything that happens with the browser open. Returns an exit code."""
    ctx.add_init_script(script=OVERLAY_JS)
    page = ctx.pages[0] if ctx.pages else ctx.new_page()
    ensure_logged_in(page, ctx)
    vid = viewer_id(ctx)
    scanned_other = state.get("viewer") and str(state["viewer"].get("id")) != str(vid)
    if scanned_other:
        print("You are logged in to a different account than last time, so the bot will scan again.")
    if args.command == "scan" or args.rescan or not state.get("scanned_at") or scanned_other:
        if not do_scan(page, ctx, state, settings, keep_txt):
            return 1
        print()
        print(f'The Excel report is ready: "{REPORT_FILE.name}"')
        print('Open it, check the "Not following back" sheet, and type yes in the Keep column for anyone you want to keep.')
        print("Save and close Excel before you start unfollowing.")
        if args.command == "scan":
            return 0
        if not ask_yes("Start unfollowing now? (Say no if you want to review the Excel first)", args.yes):
            return 0
        sync_keep_from_report(state)
    reason = do_run(page, ctx, state, settings, keep_txt, args)
    print()
    print(STOP_REASONS.get(reason, "Not started."))
    print(f'Excel report: "{REPORT_FILE}"')
    return 0 if reason in ("done", "daily_limit", "run_limit", "stopped") else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Unfollow Instagram accounts that don't follow you back, by clicking like a person.")
    ap.add_argument("command", nargs="?", default="run", choices=["scan", "run", "report"])
    ap.add_argument("--limit", type=int, default=0, help="stop after this many unfollows in this run")
    ap.add_argument("--rescan", action="store_true", help="scan your lists again before unfollowing")
    ap.add_argument("--yes", action="store_true", help="don't ask for confirmation")
    args = ap.parse_args(argv)

    settings = load_settings()
    state = load_state()
    keep_txt = read_keep_txt()
    sync_keep_from_report(state)

    if args.command == "report":
        if write_report(state, settings, keep_txt):
            print(f'Excel report updated: "{REPORT_FILE}"')
        return 0

    code = 1
    try:
        with sync_playwright() as pw:
            ctx = launch_browser(pw, settings)
            try:
                code = session(ctx, args, state, settings, keep_txt)
            finally:
                try:
                    ctx.close()
                except Exception:  # noqa: BLE001
                    pass
    except KeyboardInterrupt:
        save_state(state)
        write_report(state, settings, keep_txt)
        print("\nStopped. Progress is saved.")
        return 0
    except PlaywrightError as e:
        save_state(state)
        write_report(state, settings, keep_txt)
        if is_closed_error(e):
            print("\nThe Chrome window was closed. Progress is saved.")
            return 0
        print("\nBrowser error: " + str(e).strip().splitlines()[0][:200])
        return 1
    return code


if __name__ == "__main__":
    sys.exit(main())
